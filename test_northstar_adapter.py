#!/usr/bin/env python3
"""Test suite for Northstar adapter."""

import json
import unittest
from pathlib import Path
from northstar_adapter import NorthstarAdapter, load_json


class TestNorthstarAdapter(unittest.TestCase):
    """Test Northstar JSON → Template slide data conversion."""

    def setUp(self):
        """Load example Northstar data."""
        self.example_path = Path(__file__).parent / 'artifacts' / 'example-northstar-array.json'
        self.northstar_data = load_json(self.example_path)
        self.adapter = NorthstarAdapter()

    def test_example_data_exists(self):
        """Verify example data file exists and loads."""
        self.assertTrue(self.example_path.exists(), "Example data file missing")
        self.assertIsInstance(self.northstar_data, list)
        self.assertEqual(len(self.northstar_data), 5)

        # Verify all slide kinds present
        kinds = {slide['kind'] for slide in self.northstar_data}
        self.assertEqual(kinds, {'summary', 'performance', 'economics', 'pipeline', 'pilot'})

    def test_full_conversion(self):
        """Test complete conversion with example data."""
        performance, cost, readiness, pipeline, decisions = self.adapter.convert(self.northstar_data)

        self.assertIsNotNone(performance)
        self.assertIsNotNone(cost)
        self.assertIsNotNone(readiness)
        self.assertIsNotNone(pipeline)
        self.assertIsNotNone(decisions)
        self.assertEqual(len(self.adapter.errors), 0)

    def test_performance_conversion(self):
        """Test summary → performance slide conversion."""
        performance, cost, readiness, pipeline, decisions = self.adapter.convert(self.northstar_data)

        # Verify metrics
        self.assertEqual(len(performance.metrics), 3)
        self.assertEqual(performance.metrics[0][0], "70%")
        self.assertEqual(performance.metrics[0][1], "Tier-1 volume handled")

        # Verify narrative fields (lead AND body mapped)
        self.assertIn("97,488", performance.statement)
        self.assertIn("789,563", performance.body)

        # Verify speaker notes separate from footer
        self.assertIn("Automation rose from", performance.notes)
        self.assertIn("Source: supplied", performance.footer)

    def test_cost_conversion(self):
        """Test performance → cost slide conversion."""
        performance, cost, readiness, pipeline, decisions = self.adapter.convert(self.northstar_data)

        # Verify chart data
        self.assertIsNotNone(cost.chart_data)
        self.assertEqual(len(cost.chart_data['categories']), 4)
        self.assertEqual(cost.chart_data['values'][0], 65.8)

        # Verify table
        self.assertEqual(len(cost.table_headers), 3)
        self.assertGreater(len(cost.table_rows), 0)
        self.assertEqual(cost.table_rows[0][0], "Model requests")

        # Verify speaker notes
        self.assertIn("Request volume grew", cost.notes)

    def test_readiness_conversion(self):
        """Test economics → readiness slide conversion."""
        performance, cost, readiness, pipeline, decisions = self.adapter.convert(self.northstar_data)

        # Verify table
        self.assertEqual(len(readiness.table_headers), 4)
        self.assertEqual(len(readiness.phases), 4)
        self.assertIn("Essentials", readiness.phases[0][0])

        # Verify note field has lead AND body
        self.assertIn("Readiness gates", readiness.note)
        self.assertIn("One live pilot", readiness.note)

        # Verify speaker notes
        self.assertIn("Essentials leads", readiness.notes)

    def test_pipeline_conversion(self):
        """Test pipeline slide conversion (5th slide)."""
        performance, cost, readiness, pipeline, decisions = self.adapter.convert(self.northstar_data)

        # Verify brands
        self.assertEqual(len(pipeline.rows), 4)
        self.assertEqual(pipeline.rows[0][0], "Essentials")
        self.assertEqual(pipeline.rows[0][3], "168.5%")

        # Verify speaker notes
        self.assertIn("pipeline analysis", pipeline.notes)

    def test_decisions_conversion(self):
        """Test pilot → decisions slide conversion."""
        performance, cost, readiness, pipeline, decisions = self.adapter.convert(self.northstar_data)

        # Verify decision boxes
        self.assertEqual(len(decisions.decision_boxes), 3)
        self.assertIn("Flagship assessment", decisions.decision_boxes[0][0])

        # Verify timeline with commitment qualifications
        self.assertEqual(len(decisions.timeline), 3)
        self.assertIn("Week 1", decisions.timeline[0][0])
        self.assertIn("Status:", decisions.timeline[0][1])
        self.assertIn("Deadline:", decisions.timeline[0][1])
        self.assertIn("Consequence:", decisions.timeline[0][1])

        # Verify speaker notes include fullGates
        self.assertIn("Full gates:", decisions.notes)
        self.assertIn("Policy owner signoff", decisions.notes)

    def test_missing_slides_error(self):
        """Test error handling for missing required slides."""
        # Remove summary slide
        bad_data = [s for s in self.northstar_data if s['kind'] != 'summary']
        self.adapter.convert(bad_data)

        self.assertGreater(len(self.adapter.errors), 0)
        self.assertIn("summary", str(self.adapter.errors).lower())

    def test_chart_data_optional(self):
        """Test that chart_data can be missing."""
        # Remove chart from performance slide
        data_no_chart = []
        for slide in self.northstar_data:
            if slide['kind'] == 'performance':
                slide_copy = slide.copy()
                slide_copy.pop('chart', None)
                data_no_chart.append(slide_copy)
            else:
                data_no_chart.append(slide)

        performance, cost, readiness, pipeline, decisions = self.adapter.convert(data_no_chart)
        self.assertEqual(len(self.adapter.errors), 0)
        # Chart data should still work, just without values
        self.assertIsNotNone(cost)


class TestEndToEnd(unittest.TestCase):
    """End-to-end integration tests."""

    def test_generate_deck(self):
        """Test full pipeline: JSON → Adapter → PowerPoint."""
        from ebr_generator_clean import CleanEBRDeck

        example_path = Path(__file__).parent / 'artifacts' / 'example-northstar-array.json'
        output_path = Path(__file__).parent / 'test_northstar_output.pptx'

        # Clean up any previous test output
        if output_path.exists():
            output_path.unlink()

        try:
            # Load and convert
            northstar_data = load_json(example_path)
            adapter = NorthstarAdapter()
            performance, cost, readiness, pipeline, decisions = adapter.convert(northstar_data)

            self.assertEqual(len(adapter.errors), 0, f"Conversion failed: {adapter.errors}")

            # Generate deck
            deck = CleanEBRDeck()
            deck.build(
                performance=performance,
                cost=cost,
                readiness=readiness,
                pipeline=pipeline,
                decisions=decisions
            )
            deck.save(str(output_path))

            # Verify file created
            self.assertTrue(output_path.exists(), "Output file not created")
            self.assertGreater(output_path.stat().st_size, 1000, "Output file too small")

        finally:
            # Clean up
            if output_path.exists():
                output_path.unlink()

    def test_slide_count(self):
        """Test that output has exactly 5 slides."""
        from pptx import Presentation
        from ebr_generator_clean import CleanEBRDeck

        example_path = Path(__file__).parent / 'artifacts' / 'example-northstar-array.json'
        output_path = Path(__file__).parent / 'test_northstar_5slides.pptx'

        if output_path.exists():
            output_path.unlink()

        try:
            northstar_data = load_json(example_path)
            adapter = NorthstarAdapter()
            performance, cost, readiness, pipeline, decisions = adapter.convert(northstar_data)

            deck = CleanEBRDeck()
            deck.build(
                performance=performance,
                cost=cost,
                readiness=readiness,
                pipeline=pipeline,
                decisions=decisions
            )
            deck.save(str(output_path))

            # Load and verify slide count
            prs = Presentation(str(output_path))
            self.assertEqual(len(prs.slides), 5, "Expected exactly 5 slides")

        finally:
            if output_path.exists():
                output_path.unlink()

    def test_sentinel_values(self):
        """Test that unique sentinel values appear in output."""
        from pptx import Presentation
        from ebr_generator_clean import CleanEBRDeck

        example_path = Path(__file__).parent / 'artifacts' / 'example-northstar-array.json'
        output_path = Path(__file__).parent / 'test_northstar_sentinels.pptx'

        if output_path.exists():
            output_path.unlink()

        try:
            northstar_data = load_json(example_path)
            adapter = NorthstarAdapter()
            performance, cost, readiness, pipeline, decisions = adapter.convert(northstar_data)

            deck = CleanEBRDeck()
            deck.build(
                performance=performance,
                cost=cost,
                readiness=readiness,
                pipeline=pipeline,
                decisions=decisions
            )
            deck.save(str(output_path))

            # Load presentation and extract all text
            prs = Presentation(str(output_path))
            all_text = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        all_text.append(shape.text)
                # Check speaker notes
                if slide.has_notes_slide:
                    notes_text = slide.notes_slide.notes_text_frame.text
                    all_text.append(notes_text)

            full_text = ' '.join(all_text)

            # Test unique values per slide
            self.assertIn("168.5%", full_text, "Pipeline evidence missing")
            self.assertIn("97,488", full_text, "Performance statement missing")
            self.assertIn("789,563", full_text, "Performance body missing")

        finally:
            if output_path.exists():
                output_path.unlink()


if __name__ == '__main__':
    unittest.main()
